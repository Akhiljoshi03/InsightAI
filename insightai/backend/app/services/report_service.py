"""Generate executive report content (via OpenAI) and export to PDF/XLSX."""
from __future__ import annotations

import json
import os

import openpyxl
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from app.services import ai_service


REPORT_SYSTEM_PROMPT = """You are a senior business analyst writing an executive report
from a dataset profile. Respond ONLY as JSON:
{"executive_summary": "...", "dataset_overview": "...", "insights": ["...", "..."],
 "opportunities": ["...", "..."], "risks": ["...", "..."],
 "recommendations": ["...", "..."], "predictions": "...", "conclusion": "..."}
Keep each string field to 2-4 sentences and each list to 3-6 items."""


def generate_report_content(profile: dict, title: str) -> dict:
    client = ai_service.get_client()
    from app.config import get_settings
    settings = get_settings()
    response = client.chat.completions.create(
        model=settings.openai_model,
        messages=[
            {"role": "system", "content": REPORT_SYSTEM_PROMPT},
            {"role": "user", "content": f"Report title: {title}\nDataset profile:\n{json.dumps(profile, default=str)[:14000]}"},
        ],
        temperature=0.3,
        response_format={"type": "json_object"},
    )
    return json.loads(response.choices[0].message.content)


def export_pdf(title: str, content: dict, out_path: str) -> str:
    doc = SimpleDocTemplate(out_path, pagesize=LETTER)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=colors.HexColor("#4C1D95"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=colors.HexColor("#6D28D9"))
    body = styles["BodyText"]

    story = [Paragraph(title, h1), Spacer(1, 16)]
    section_titles = {
        "executive_summary": "Executive Summary",
        "dataset_overview": "Dataset Overview",
        "insights": "Key Insights",
        "opportunities": "Business Opportunities",
        "risks": "Risks",
        "recommendations": "Recommendations",
        "predictions": "Predictions",
        "conclusion": "Conclusion",
    }
    for key, label in section_titles.items():
        value = content.get(key)
        if not value:
            continue
        story.append(Paragraph(label, h2))
        if isinstance(value, list):
            for item in value:
                story.append(Paragraph(f"&bull; {item}", body))
        else:
            story.append(Paragraph(str(value), body))
        story.append(Spacer(1, 10))

    doc.build(story)
    return out_path


def export_xlsx(title: str, content: dict, profile: dict, out_path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)

    row = 3
    for key, label in [
        ("executive_summary", "Executive Summary"), ("dataset_overview", "Dataset Overview"),
        ("predictions", "Predictions"), ("conclusion", "Conclusion"),
    ]:
        if content.get(key):
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            row += 1
            ws.cell(row=row, column=1, value=str(content[key]))
            row += 2

    for key, label in [("insights", "Key Insights"), ("opportunities", "Opportunities"),
                        ("risks", "Risks"), ("recommendations", "Recommendations")]:
        items = content.get(key) or []
        if items:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            row += 1
            for item in items:
                ws.cell(row=row, column=1, value=f"- {item}")
                row += 1
            row += 1

    ws2 = wb.create_sheet("Dataset Profile")
    ws2.append(["Column", "Type", "Nulls %", "Unique", "Mean", "Std"])
    for col in profile.get("columns", []):
        ws2.append([col.get("name"), col.get("dtype"), col.get("null_pct"),
                    col.get("unique_count"), col.get("mean"), col.get("std")])

    wb.save(out_path)
    return out_path
